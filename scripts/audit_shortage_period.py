# -*- coding: utf-8 -*-
"""Focused audit: period filter + abs shortage + sheet totals."""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.parser import filter_by_period, parse_network_excel  # noqa: E402

SRC = Path(r"C:\Users\Администратор\Documents\ЗЯ\ПЮ\Инвет\Торговый_зал\АнализАвгуст.xlsx")
REPORT = Path(r"C:\Users\Администратор\Documents\ЗЯ\ПЮ\Инвет\Анализ_розница_АнализАвгуст_20260909_124410.xlsx")
KPI = dict(surplus=5182748.20, shortage=9971173.37, docs=1237, stores=15)


def main() -> None:
    print("Parsing full August source...")
    df, meta = parse_network_excel(str(SRC))
    sur0 = float(df["излишек_сумма"].sum())
    sh0 = float(df["недостача_сумма"].sum())
    print(f"full: rows={len(df)} docs={meta.doc_count} stores={len(meta.stores)}")
    print(f"full: surplus={sur0:,.2f} shortage={sh0:,.2f}")
    print(f"dates: {meta.date_min} .. {meta.date_max}")

    # How shortage signs look in raw parsed fields
    print(
        f"rows with both surplus and shortage >0: "
        f"{((df['излишек_сумма']>0)&(df['недостача_сумма']>0)).sum()}"
    )

    d = pd.to_datetime(df["дата_док"]).dt.date
    slices = {
        "all": d.notna(),
        "Aug1-31": (d >= date(2026, 8, 1)) & (d <= date(2026, 8, 31)),
        "Aug2-31": (d >= date(2026, 8, 2)) & (d <= date(2026, 8, 31)),
        "period30_end31": None,
    }
    dfp = filter_by_period(df, 30, date(2026, 8, 31))
    print(
        f"filter_by_period(30, 31.08): start={dfp.attrs.get('period_start')} "
        f"end={dfp.attrs.get('period_end')} rows={len(dfp)} "
        f"docs={dfp['номер_док'].nunique()} "
        f"sur={dfp['излишек_сумма'].sum():,.2f} sh={dfp['недостача_сумма'].sum():,.2f}"
    )

    for name, mask in slices.items():
        if mask is None:
            continue
        sub = df.loc[mask]
        print(
            f"{name}: rows={len(sub)} docs={sub['номер_док'].nunique()} "
            f"sur={sub['излишек_сумма'].sum():,.2f} sh={sub['недостача_сумма'].sum():,.2f}"
        )

    # Outside Aug2-31
    outside = df.loc[~((d >= date(2026, 8, 2)) & (d <= date(2026, 8, 31)))]
    print(
        f"OUTSIDE Aug2-31: rows={len(outside)} docs={outside['номер_док'].nunique()} "
        f"sur={outside['излишек_сумма'].sum():,.2f} sh={outside['недостача_сумма'].sum():,.2f}"
    )
    if len(outside):
        g = (
            outside.assign(day=d.loc[outside.index])
            .groupby("day")
            .agg(
                docs=("номер_док", "nunique"),
                sur=("излишек_сумма", "sum"),
                sh=("недостача_сумма", "sum"),
            )
        )
        print(g.to_string())

    aug = df.loc[(d >= date(2026, 8, 2)) & (d <= date(2026, 8, 31))]
    print("\nDELTA Aug2-31 vs REPORT KPI:")
    print(f"  surplus {aug['излишек_сумма'].sum() - KPI['surplus']:,.2f}")
    print(f"  shortage {aug['недостача_сумма'].sum() - KPI['shortage']:,.2f}")
    print(f"  docs {aug['номер_док'].nunique()} vs KPI {KPI['docs']}")

    # Compare report Недостачи sheet total
    wb = load_workbook(REPORT, data_only=True, read_only=True)
    for sheet in ("Недостачи", "Излишки", "Рейтинг магазинов"):
        ws = wb[sheet]
        print(f"\n--- sheet {sheet} header scan ---")
        for r in range(1, 6):
            vals = [ws.cell(r, c).value for c in range(1, 14)]
            if any(v is not None for v in vals):
                print(f"  R{r}: {vals}")
    # Рейтинг: sum недостачи column
    ws = wb["Рейтинг магазинов"]
    # find header
    hdr_r = None
    for r in range(1, 10):
        for c in range(1, 20):
            if ws.cell(r, c).value and str(ws.cell(r, c).value).strip().lower() == "недостачи":
                hdr_r, sh_c = r, c
                sur_c = None
                for c2 in range(1, 20):
                    if ws.cell(r, c2).value and "излиш" in str(ws.cell(r, c2).value).lower():
                        sur_c = c2
                break
        if hdr_r:
            break
    tot_sh = tot_sur = 0.0
    n = 0
    for r in range(hdr_r + 1, (ws.max_row or hdr_r) + 1):
        name = ws.cell(r, 2).value  # usually store name col
        shv = ws.cell(r, sh_c).value
        srv = ws.cell(r, sur_c).value if sur_c else None
        if isinstance(shv, (int, float)):
            tot_sh += float(shv)
            n += 1
        if isinstance(srv, (int, float)):
            tot_sur += float(srv)
    print(f"Рейтинг sum недостачи rows={n} sh={tot_sh:,.2f} sur={tot_sur:,.2f}")

    # Недостачи sheet: find Сумма column
    ws = wb["Недостачи"]
    for r in range(1, 8):
        rowv = [ws.cell(r, c).value for c in range(1, 12)]
        print(f"Недостачи R{r}: {rowv}")
    # assume modern sheet header at row 4
    sum_c = None
    hr = None
    for r in range(1, 8):
        for c in range(1, 15):
            v = ws.cell(r, c).value
            if v and str(v).strip().lower() == "сумма":
                hr, sum_c = r, c
    tot = 0.0
    n = 0
    if hr and sum_c:
        for r in range(hr + 1, (ws.max_row or hr) + 1):
            v = ws.cell(r, sum_c).value
            if isinstance(v, (int, float)):
                tot += float(v)
                n += 1
    print(f"Недостачи sheet sum={tot:,.2f} n={n}")
    wb.close()

    # Names containing итого that are products
    bad = df[df["наименование"].str.lower().str.contains(r"итого|^\s*всего\s*$", regex=True)]
    print(f"\nproduct names like итого/всего: {len(bad)}")
    if len(bad):
        print(bad["наименование"].value_counts().head(20).to_string())
        print(
            f"  contribution sur={bad['излишек_сумма'].sum():,.2f} "
            f"sh={bad['недостача_сумма'].sum():,.2f}"
        )

    # Document totals accidentally parsed? (same sums as doc header)
    # Check outline: products under doc should not include blank groups
    print("\nDone.")


if __name__ == "__main__":
    main()
