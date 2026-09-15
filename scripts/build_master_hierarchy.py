#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Rebuild data/master_hierarchy.json from Шаблон_иерархия.xlsx.

Usage (from project root):
  python scripts/build_master_hierarchy.py --source "C:\\path\\to\\Шаблон_иерархия.xlsx"

Runtime of the analyzer does NOT need this file — only maintainers re-extracting the etalon.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "data" / "master_hierarchy.json"
DEFAULT_CSV = ROOT / "data" / "master_hierarchy.csv"

HEADER_SKIP_PREFIXES = (
    "остатки и обороты",
    "отбор:",
    "показатели:",
    "итоги по:",
    "номенклатура",
)


def extract(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    levels = {
        int(row_idx): int(getattr(dim, "outlineLevel", 0) or 0)
        for row_idx, dim in ws.row_dimensions.items()
    }

    records = []
    stack: dict[int, str] = {}
    max_depth = 0

    for i in range(1, ws.max_row + 1):
        b = ws.cell(i, 2).value
        c = ws.cell(i, 3).value
        lv = levels.get(i, 0)
        b_s = str(b).strip() if b is not None else ""
        c_s = str(c).strip() if c is not None else ""

        if c_s:
            path_levels = {k: stack[k] for k in sorted(stack) if k < lv}
            hier = [path_levels[k] for k in sorted(path_levels)]
            max_depth = max(max_depth, len(hier))
            rec = {
                "name": c_s,
                "leaf_group": hier[-1] if hier else "",
                "hierarchy_path": hier,
                "outline_level": lv,
            }
            for idx, node in enumerate(hier):
                rec[f"level_{idx}"] = node
            records.append(rec)
            continue

        if not b_s:
            continue
        low = b_s.lower()
        if any(low.startswith(p) or low == p for p in HEADER_SKIP_PREFIXES):
            continue
        for k in list(stack.keys()):
            if k >= lv:
                del stack[k]
        stack[lv] = b_s

    wb.close()

    seen = set()
    items = []
    for r in records:
        if r["name"] in seen:
            continue
        seen.add(r["name"])
        items.append(r)

    return {
        "source_file": path.name,
        "source_sheet": "Лист1",
        "extracted_rows": len(records),
        "unique_names_stored": len(items),
        "max_hierarchy_depth": max_depth,
        "level_names": [f"level_{i}" for i in range(max_depth)],
        "notes": [
            "Иерархия из Excel outline: колонка B = узлы, колонка C = SKU.",
            "level_0..level_N — путь от корня к родителю SKU.",
            "leaf_group — ближайший родитель (category_group для пересорта).",
        ],
        "top_level_counts": dict(Counter(it.get("level_0", "") for it in items)),
        "items": items,
    }


def main() -> None:
    ap = argparse.ArgumentParser(description="Rebuild embedded master hierarchy")
    ap.add_argument("--source", required=True, help="Path to Шаблон_иерархия.xlsx")
    ap.add_argument("--out-json", default=str(DEFAULT_OUT))
    ap.add_argument("--out-csv", default=str(DEFAULT_CSV))
    args = ap.parse_args()

    src = Path(args.source)
    if not src.exists():
        print(f"Source not found: {src}", file=sys.stderr)
        sys.exit(1)

    payload = extract(src)
    out_json = Path(args.out_json)
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    out_csv = Path(args.out_csv)
    depth = payload["max_hierarchy_depth"]
    fields = ["name", "leaf_group", "hierarchy_path"] + [f"level_{i}" for i in range(depth)] + ["outline_level"]
    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for it in payload["items"]:
            row = dict(it)
            row["hierarchy_path"] = " / ".join(it["hierarchy_path"])
            w.writerow(row)

    print(f"OK items={payload['unique_names_stored']} depth={depth}")
    print(f"JSON: {out_json}")
    print(f"CSV:  {out_csv}")


if __name__ == "__main__":
    main()
