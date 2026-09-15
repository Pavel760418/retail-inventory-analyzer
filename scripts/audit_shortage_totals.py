# -*- coding: utf-8 -*-
"""Full audit: report shortage/surplus vs source TDSheet."""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT_PKG = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT_PKG))

from src.parser import (  # noqa: E402
    DOC_RE,
    _read_outline_levels,
    _to_float,
    parse_network_excel,
)

ROOT = Path(r"C:\Users\Администратор\Documents\ЗЯ\ПЮ\Инвет")
REPORT = ROOT / "Анализ_розница_АнализАвгуст_20260909_124410.xlsx"
SRC = ROOT / "Торговый_зал" / "АнализАвгуст.xlsx"
BLOCKS = [
    ROOT / "Торговый_зал" / "АнализАвгустЭкспресс.xlsx",
    ROOT / "Торговый_зал" / "АнализАвгуст1блок.xlsx",
    ROOT / "Торговый_зал" / "АнализАвгуст2блок.xlsx",
]
REPORT_KPI = {"surplus": 5182748.20, "shortage": 9971173.37}


def source_totals_scan(path: Path) -> None:
    xl = pd.ExcelFile(path)
    sheet = xl.sheet_names[0]
    df = pd.read_excel(path, sheet_name=sheet, header=None)
    levels = _read_outline_levels(str(path), sheet)

    print(f"\n=== SOURCE: {path.name} sheet={sheet} rows={len(df)} ===")
    for i in range(0, 8):
        a = df.iloc[i, 0]
        nums = [df.iloc[i, k] for k in range(1, 9)]
        print(f"  raw[{i}]: {a!r} | cols1-8={nums}")

    itogo = []
    store_headers = []
    for i in range(len(df)):
        a_val = df.iloc[i, 0]
        if pd.isna(a_val):
            continue
        a = str(a_val).strip()
        if not a:
            continue
        low = a.lower()
        lv = levels.get(i + 1, 0)
        vals = [_to_float(df.iloc[i, k]) for k in range(1, 9)]
        # col7 = vals[6] сумма разница? col8=vals[7] — уточним по шапке
        if "итого" in low or low in {"всего", "итог"}:
            itogo.append((i + 1, lv, a, vals[6], vals[7], vals))
        if lv == 0 and not DOC_RE.match(a) and i >= 3:
            store_headers.append((i + 1, a, vals[6], vals[7], vals))

    print(f"  Итого-like rows: {len(itogo)}")
    for t in itogo[:25]:
        print(
            f"    ExcelR{t[0]} lv={t[1]} c7={t[3]:.2f} c8={t[4]:.2f} "
            f"full={t[5]} name={t[2][:90]}"
        )
    if len(itogo) > 25:
        print(f"    ... +{len(itogo) - 25} more")
    print(
        f"  SUM ALL Итого-like c7={sum(t[3] for t in itogo):,.2f} "
        f"c8={sum(t[4] for t in itogo):,.2f}"
    )

    stores_only = [s for s in store_headers if "итого" not in s[1].lower()]
    print(f"  Store headers w/o итого: {len(stores_only)}")
    # In 1C export: typically col7=surplus sum, col8=shortage sum (need verify header)
    print(
        f"  SUM store headers c7={sum(s[2] for s in stores_only):,.2f} "
        f"c8={sum(s[3] for s in stores_only):,.2f}"
    )
    for s in stores_only:
        print(
            f"    R{s[0]} {s[1][:55]:55s} c7={s[2]:>14,.2f} c8={s[3]:>14,.2f} "
            f"vals={s[4]}"
        )


def parse_and_sum(path: Path):
    df, meta = parse_network_excel(str(path))
    sur = float(df["излишек_сумма"].sum())
    sh = float(df["недостача_сумма"].sum())
    nt = meta.network_total_row
    print(f"\n=== PARSER: {path.name} ===")
    print(f"  rows={len(df)} stores={len(meta.stores)} docs={meta.doc_count}")
    print(f"  parsed surplus={sur:,.2f} shortage={sh:,.2f}")
    if nt:
        print(f"  network_total name={nt.get('магазин')!r}")
        for k in sorted(nt.keys()):
            if k.startswith("c"):
                print(f"    {k}={nt[k]}")
    bad = df[df["наименование"].str.lower().str.contains("итого|всего", regex=True)]
    print(
        f"  products with итого/всего: {len(bad)} "
        f"sur={bad['излишек_сумма'].sum():,.2f} sh={bad['недостача_сумма'].sum():,.2f}"
    )
    if len(bad):
        print(
            bad[["магазин", "наименование", "излишек_сумма", "недостача_сумма"]]
            .head(15)
            .to_string()
        )

    # Duplicate product lines? same store+doc+name
    dup_key = df.groupby(["магазин", "документ", "наименование"]).size()
    dups = dup_key[dup_key > 1]
    print(f"  duplicate store+doc+name groups: {len(dups)}")
    if len(dups):
        print(dups.head(10))

    return df, meta, sur, sh


def sheet_sum(report: Path, sheet: str, col_name_substr: str) -> float:
    wb = load_workbook(report, data_only=True, read_only=True)
    ws = wb[sheet]
    # find header row
    header_row = None
    col_idx = None
    for r in range(1, 10):
        for c in range(1, 20):
            v = ws.cell(r, c).value
            if v and col_name_substr.lower() in str(v).lower():
                header_row = r
                col_idx = c
                break
        if header_row:
            break
    total = 0.0
    n = 0
    if header_row and col_idx:
        for r in range(header_row + 1, (ws.max_row or header_row) + 1):
            v = ws.cell(r, col_idx).value
            if isinstance(v, (int, float)):
                total += float(v)
                n += 1
    wb.close()
    print(f"  sheet '{sheet}' col~'{col_name_substr}' rows={n} sum={total:,.2f}")
    return total


def main() -> None:
    print("REPORT KPI:", REPORT_KPI)
    source_totals_scan(SRC)
    df, meta, sur, sh = parse_and_sum(SRC)

    print("\n=== DELTA parser vs report KPI ===")
    print(f"  surplus delta = {sur - REPORT_KPI['surplus']:,.2f}")
    print(f"  shortage delta = {sh - REPORT_KPI['shortage']:,.2f}")

    print("\n=== BLOCKS sum ===")
    block_sur = block_sh = 0.0
    for b in BLOCKS:
        if not b.exists():
            print(f"  missing {b}")
            continue
        _, _, bs, bh = parse_and_sum(b)
        block_sur += bs
        block_sh += bh
    print(f"  BLOCKS TOTAL surplus={block_sur:,.2f} shortage={block_sh:,.2f}")
    print(f"  vs merged file delta sur={block_sur - sur:,.2f} sh={block_sh - sh:,.2f}")

    print("\n=== Report sheets Недостачи / Излишки ===")
    sheet_sum(REPORT, "Недостачи", "Сумма")
    sheet_sum(REPORT, "Излишки", "Сумма")

    # Store-level compare: parser store sum vs store header row in source
    print("\n=== Per-store: parser sum vs source outline0 header ===")
    xl = pd.ExcelFile(SRC)
    sheet = xl.sheet_names[0]
    raw = pd.read_excel(SRC, sheet_name=sheet, header=None)
    levels = _read_outline_levels(str(SRC), sheet)
    header_by_store = {}
    for i in range(len(raw)):
        a_val = raw.iloc[i, 0]
        if pd.isna(a_val):
            continue
        a = str(a_val).strip()
        lv = levels.get(i + 1, 0)
        if lv == 0 and a and not DOC_RE.match(a) and "итого" not in a.lower() and i >= 4:
            # In parser: col7=vals[7] is surplus? parser uses vals[7] for surplus_sum max, vals[8] shortage
            # vals dict is 1..8 from columns
            c7 = _to_float(raw.iloc[i, 7])  # pandas col index 7 = Excel col H = сумма излишек?
            c8 = _to_float(raw.iloc[i, 8])
            header_by_store[a] = (c7, c8)

    by_store = (
        df.groupby("магазин")
        .agg(sur=("излишек_сумма", "sum"), sh=("недостача_сумма", "sum"))
        .reset_index()
    )
    print(f"{'магазин':40s} {'hdr_sur':>14} {'par_sur':>14} {'d_sur':>12} {'hdr_sh':>14} {'par_sh':>14} {'d_sh':>12}")
    for _, row in by_store.iterrows():
        m = row["магазин"]
        hs, hh = header_by_store.get(m, (None, None))
        if hs is None:
            print(f"{m[:40]:40s} {'N/A':>14} {row['sur']:14,.2f} {'':>12} {'N/A':>14} {row['sh']:14,.2f}")
            continue
        print(
            f"{m[:40]:40s} {hs:14,.2f} {row['sur']:14,.2f} {row['sur']-hs:12,.2f} "
            f"{hh:14,.2f} {row['sh']:14,.2f} {row['sh']-hh:12,.2f}"
        )

    # Check column headers in source
    print("\n=== Source column headers (row looking for markers) ===")
    for i in range(min(6, len(raw))):
        print(i, [raw.iloc[i, k] for k in range(0, 9)])


if __name__ == "__main__":
    main()
