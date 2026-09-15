# -*- coding: utf-8 -*-
"""Analysis orchestration for Streamlit — wraps existing business logic, no recalc."""
from __future__ import annotations

import datetime
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from src.anomalies import anomalies_summary, detect_book_sum_anomalies
from src.author import (
    annotate_overlap_authors,
    author_network_summary,
    calc_author_role_stats,
    calc_author_store_stats,
    detect_author_chains,
    enrich_author_columns,
    enrich_store_metrics_with_authors,
)
from src.capitalization import (
    capitalization_summary,
    match_capitalization_to_inventory,
    parse_capitalization_excel,
)
from src.catalog import enrich_dataframe, load_catalog, unmatched_summary
from src.excel.export import build_report
from src.metrics import (
    calc_network_summary,
    calc_sku_cross,
    calc_store_metrics,
    generate_conclusions,
)
from src.models import Config, ParseMeta
from src.overlap import build_analytical_cross_store, build_operational_overlaps
from src.parser import filter_by_period, parse_network_excel


@dataclass
class QualityReport:
    loaded_rows: int = 0
    after_period_rows: int = 0
    excluded_by_period: int = 0
    stores: int = 0
    docs: int = 0
    hierarchy_unmatched_rows: int = 0
    hierarchy_unmatched_names: int = 0
    anomaly_total: int = 0
    anomaly_critical: int = 0
    empty_name_rows: int = 0
    duplicate_keys: int = 0
    passed: bool = True
    issues: List[Dict[str, str]] = field(default_factory=list)


@dataclass
class AnalysisResult:
    df: pd.DataFrame
    meta: ParseMeta
    summary: Dict[str, Any]
    store_metrics: pd.DataFrame
    sku_cross: pd.DataFrame
    chronic: pd.DataFrame
    overlap_op: pd.DataFrame
    anom_df: pd.DataFrame
    anom_sum: Dict[str, Any]
    author_role: pd.DataFrame
    author_store: pd.DataFrame
    author_chains: pd.DataFrame
    author_sum: Dict[str, Any]
    cap_matched: pd.DataFrame
    cap_sum: Dict[str, Any]
    conclusions: List[Tuple[str, str]]
    quality: QualityReport
    excel_bytes: bytes
    excel_sheets: int
    excel_seconds: float
    period_str: str
    source_fingerprint: str


def _fingerprint(*names: str) -> str:
    return "|".join(names)


def _build_quality(df_raw_len: int, df: pd.DataFrame, um: pd.DataFrame, anom_sum: Dict) -> QualityReport:
    q = QualityReport()
    q.loaded_rows = int(df_raw_len)
    q.after_period_rows = int(len(df))
    q.excluded_by_period = max(df_raw_len - len(df), 0)
    q.stores = int(df["магазин"].nunique()) if not df.empty else 0
    q.docs = int(df["документ"].nunique()) if not df.empty else 0
    q.hierarchy_unmatched_rows = int((~df["hierarchy_matched"]).sum()) if "hierarchy_matched" in df.columns else 0
    q.hierarchy_unmatched_names = int(len(um)) if um is not None else 0
    q.anomaly_total = int(anom_sum.get("total", 0))
    q.anomaly_critical = int(anom_sum.get("critical", 0))
    q.empty_name_rows = int(df["наименование"].isna().sum()) if "наименование" in df.columns else 0
    if {"магазин", "документ", "наименование"}.issubset(df.columns):
        q.duplicate_keys = int(df.duplicated(subset=["магазин", "документ", "наименование"]).sum())

    if q.empty_name_rows:
        q.issues.append({
            "причина": "Пустые наименования",
            "количество": str(q.empty_name_rows),
            "исправление": "Проверьте строки номенклатуры в исходной выгрузке 1С.",
        })
    if q.hierarchy_unmatched_rows:
        q.issues.append({
            "причина": "SKU вне эталонной иерархии",
            "количество": f"{q.hierarchy_unmatched_rows} строк / {q.hierarchy_unmatched_names} имён",
            "исправление": "Обновите шаблон иерархии или сверьте наименование с справочником.",
        })
    if q.anomaly_critical:
        q.issues.append({
            "причина": "Критические аномалии книжных сумм (тип 2)",
            "количество": str(q.anomaly_critical),
            "исправление": "Заполните книжную нормативную и фактическую суммы в 1С.",
        })
    # Soft pass: structural parse succeeded; issues are informational
    q.passed = q.empty_name_rows == 0
    return q


def run_full_analysis(
    inventory_path: str,
    capitalization_path: str,
    *,
    period_days: int = 30,
    end_date: Optional[datetime.date] = None,
    enable_cross_store: bool = False,
    source_names: Optional[Tuple[str, str]] = None,
) -> AnalysisResult:
    """Run existing pipeline steps and produce UI metrics + Excel bytes."""
    df, meta = parse_network_excel(inventory_path)
    raw_len = int(meta.sku_count or len(df))
    if "автор_дока" not in df.columns:
        df = enrich_author_columns(df)

    end = end_date or meta.date_max
    df = filter_by_period(df, period_days, end)
    if df.empty:
        raise ValueError("После фильтрации по периоду не осталось данных. Увеличьте период или проверьте даты документов.")

    p_start = df.attrs.get("period_start", meta.date_min)
    p_end = df.attrs.get("period_end", meta.date_max)
    period_str = f"{p_start.strftime('%d.%m.%Y')} — {p_end.strftime('%d.%m.%Y')}"

    df = enrich_dataframe(df, load_catalog(None))
    um = unmatched_summary(df)

    overlap_op = build_operational_overlaps(df)
    overlap_op = annotate_overlap_authors(overlap_op, df)
    if enable_cross_store:
        _ = build_analytical_cross_store(df)

    store_metrics = calc_store_metrics(df, overlap_op)
    sku_cross = calc_sku_cross(df)
    chronic = sku_cross[sku_cross["chronic"]]
    summary = calc_network_summary(df, overlap_op)

    anom_df = detect_book_sum_anomalies(df)
    anom_sum = anomalies_summary(anom_df, total_rows=len(df))
    author_role = calc_author_role_stats(df, overlap_op)
    author_store = calc_author_store_stats(df, overlap_op)
    author_chains = detect_author_chains(df)
    author_sum = author_network_summary(author_role, author_store, author_chains)
    store_metrics = enrich_store_metrics_with_authors(store_metrics, author_store)

    cap_raw = parse_capitalization_excel(capitalization_path)
    cap_matched = match_capitalization_to_inventory(cap_raw, df)
    cap_sum = capitalization_summary(cap_matched)

    conclusions = generate_conclusions(
        summary, store_metrics, chronic,
        cap_summary=cap_sum,
        anom_summary=anom_sum,
        author_summary=author_sum,
    )
    quality = _build_quality(raw_len, df, um, anom_sum)

    names = source_names or (meta.file_name, Path(capitalization_path).name)
    fp = _fingerprint(*names, period_str, str(period_days))

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    out_path = tmp.name
    tmp.close()
    t0 = time.perf_counter()
    try:
        cfg = Config(
            input_path=inventory_path,
            output_path=out_path,
            period_days=period_days,
            end_date=end_date,
            enable_cross_store=enable_cross_store,
            capitalization_path=capitalization_path,
        )
        build_report(cfg)
        excel_bytes = Path(out_path).read_bytes()
        from openpyxl import load_workbook
        wb = load_workbook(out_path, read_only=True)
        n_sheets = len(wb.sheetnames)
        wb.close()
    finally:
        try:
            Path(out_path).unlink(missing_ok=True)
        except OSError:
            pass
    elapsed = time.perf_counter() - t0

    return AnalysisResult(
        df=df,
        meta=meta,
        summary=summary,
        store_metrics=store_metrics,
        sku_cross=sku_cross,
        chronic=chronic,
        overlap_op=overlap_op,
        anom_df=anom_df,
        anom_sum=anom_sum,
        author_role=author_role,
        author_store=author_store,
        author_chains=author_chains,
        author_sum=author_sum,
        cap_matched=cap_matched,
        cap_sum=cap_sum,
        conclusions=conclusions,
        quality=quality,
        excel_bytes=excel_bytes,
        excel_sheets=n_sheets,
        excel_seconds=elapsed,
        period_str=period_str,
        source_fingerprint=fp,
    )
