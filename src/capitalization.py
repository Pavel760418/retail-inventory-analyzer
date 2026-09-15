# -*- coding: utf-8 -*-
"""Parse and match shift-close surplus capitalization (оприходование излишков).

Rules (Release 4):
- exact name matching only (strip + normalized sku_key equality);
- never mix warehouses / stores;
- warehouse→store mapping must be unique; ambiguous → uncertain.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Optional, Set, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from zipfile import BadZipFile

from src.text_normalize import norm_txt, sku_key

_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _detect_excel_container(path: Path) -> str:
    """Return ``xlsx``, ``xls`` (BIFF/OLE), or ``unknown`` by file header (not extension)."""
    with path.open("rb") as fh:
        head = fh.read(8)
    if head[:2] == b"PK":
        return "xlsx"
    if head[:8] == _OLE_SIGNATURE:
        return "xls"
    return "unknown"


def _capitalization_format_error(path: Path) -> ValueError:
    ext = path.suffix.lower()
    hint = (
        "Файл повреждён или это не Excel. Откройте его в Excel и сохраните как «Книга Excel (.xlsx)»."
    )
    if ext == ".xlsx":
        kind = _detect_excel_container(path)
        if kind == "xls":
            hint = (
                "Файл имеет расширение .xlsx, но внутри это старый формат Excel (.xls). "
                "Откройте в Excel и сохраните как «Книга Excel (.xlsx)» "
                "или оставьте расширение .xls — программа теперь читает оба формата."
            )
    return ValueError(f"Не удалось прочитать файл оприходования «{path.name}». {hint}")

STATUS_MATCHED = "matched"
STATUS_MATCHED_SHORTAGE = "matched_with_shortage"
STATUS_UNMATCHED_NAME = "unmatched_name"
STATUS_UNMATCHED_STORE = "unmatched_store"
STATUS_UNCERTAIN_STORE = "uncertain_store"

_WAREHOUSE_PREFIXES = ("склад", "рц", "фабрика")
_STOP_TOKENS = {
    "склад", "компании", "магазин", "производство", "зеленого", "яблока",
    "и", "в", "на", "от", "по", "г", "ул",
}


def _to_float(v) -> float:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _is_warehouse_header(name: str) -> bool:
    low = name.lower().strip()
    return any(low.startswith(p) for p in _WAREHOUSE_PREFIXES)


def _find_cap_header(
    cell_text: Callable[[int, int], str],
    max_row: int,
) -> Tuple[Optional[int], int, int]:
    """Locate header row and qty/sum column indices (1-based columns)."""
    header_row = None
    qty_col = 4
    sum_col = 6
    for i in range(1, min(30, max_row + 1)):
        vals = [cell_text(i, c).lower() for c in range(1, 9)]
        joined = " | ".join(vals)
        if "закрытие смены количество" in joined and "закрытие смены сумма" in joined:
            header_row = i
            for c in range(1, 9):
                v = cell_text(i, c).lower()
                if "закрытие смены количество" in v:
                    qty_col = c
                if "закрытие смены сумма" in v:
                    sum_col = c
            break
    return header_row, qty_col, sum_col


def _parse_cap_rows_from_sheet(
    *,
    max_row: int,
    cell_value,
    outline_level,
    header_row: int,
    qty_col: int,
    sum_col: int,
) -> List[dict]:
    rows: List[dict] = []
    current_wh = ""
    for i in range(header_row + 1, max_row + 1):
        name_raw = cell_value(i, 1)
        if name_raw is None or str(name_raw).strip() == "":
            continue
        name = str(name_raw).strip()
        low = name.lower()
        if low == "номенклатура":
            continue
        # Пропускаем подитоги/итоги выгрузки 1С — иначе сумма оприходования удваивается
        if low in {"итого", "всего", "итог"} or low.startswith("итого ") or low.startswith("всего "):
            continue
        lv = outline_level(i)
        qty = _to_float(cell_value(i, qty_col))
        sm = _to_float(cell_value(i, sum_col))

        if lv == 0 and _is_warehouse_header(name):
            current_wh = name
            continue
        if _is_warehouse_header(name) and lv == 0:
            current_wh = name
            continue
        if not current_wh:
            rows.append({
                "склад": "",
                "наименование": name,
                "оп_количество": qty,
                "оп_сумма": sm,
                "sku_key": sku_key(name),
                "name_exact": name,
            })
            continue
        if lv >= 1 or not _is_warehouse_header(name):
            if _is_warehouse_header(name):
                current_wh = name
                continue
            rows.append({
                "склад": current_wh,
                "наименование": name,
                "оп_количество": qty,
                "оп_сумма": sm,
                "sku_key": sku_key(name),
                "name_exact": name,
            })
    return rows


def _parse_capitalization_xlsx(filepath: str, sheet_name: Optional[str]) -> List[dict]:
    try:
        wb = load_workbook(filepath, read_only=False, data_only=True)
    except BadZipFile as exc:
        raise _capitalization_format_error(Path(filepath)) from exc
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    levels = {
        int(idx): int(getattr(dim, "outlineLevel", 0) or 0)
        for idx, dim in ws.row_dimensions.items()
    }

    def cell_value(row: int, col: int):
        return ws.cell(row, col).value

    def cell_text(row: int, col: int) -> str:
        return str(cell_value(row, col) or "")

    header_row, qty_col, sum_col = _find_cap_header(cell_text, ws.max_row)
    if header_row is None:
        wb.close()
        raise ValueError(
            "Файл оприходования не распознан: нет колонок "
            "«Закрытие смены количество» / «Закрытие смены сумма»."
        )
    rows = _parse_cap_rows_from_sheet(
        max_row=ws.max_row,
        cell_value=cell_value,
        outline_level=lambda r: levels.get(r, 0),
        header_row=header_row,
        qty_col=qty_col,
        sum_col=sum_col,
    )
    wb.close()
    return rows


def _parse_capitalization_xls(filepath: str, sheet_name: Optional[str]) -> List[dict]:
    try:
        import xlrd
    except ImportError as exc:
        raise ImportError(
            "Для файлов Excel старого формата (.xls) нужен пакет xlrd. "
            "Выполните: pip install xlrd==1.2.0"
        ) from exc

    book = xlrd.open_workbook(filepath)
    if sheet_name:
        sh = book.sheet_by_name(sheet_name)
    else:
        sh = book.sheet_by_index(0)

    def cell_value(row: int, col: int):
        # xlrd is 0-based; API here is 1-based like openpyxl
        r, c = row - 1, col - 1
        if r < 0 or c < 0 or r >= sh.nrows or c >= sh.ncols:
            return None
        return sh.cell_value(r, c)

    def cell_text(row: int, col: int) -> str:
        return str(cell_value(row, col) or "")

    header_row, qty_col, sum_col = _find_cap_header(cell_text, sh.nrows)
    if header_row is None:
        raise ValueError(
            "Файл оприходования не распознан: нет колонок "
            "«Закрытие смены количество» / «Закрытие смены сумма»."
        )
    return _parse_cap_rows_from_sheet(
        max_row=sh.nrows,
        cell_value=cell_value,
        outline_level=lambda _r: 0,
        header_row=header_row,
        qty_col=qty_col,
        sum_col=sum_col,
    )


def parse_capitalization_excel(filepath: str, sheet_name: Optional[str] = None) -> pd.DataFrame:
    """Parse TDSheet-like capitalization export into flat rows per warehouse+SKU.

    Expected columns (row with headers):
    - Склад компании / Номенклатура in col A
    - Закрытие смены количество
    - Закрытие смены сумма
    """
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"Файл оприходования не найден: {filepath}")

    kind = _detect_excel_container(path)
    if kind == "xlsx":
        rows = _parse_capitalization_xlsx(filepath, sheet_name)
    elif kind == "xls":
        rows = _parse_capitalization_xls(filepath, sheet_name)
    else:
        raise _capitalization_format_error(path)

    if not rows:
        raise ValueError("В файле оприходования не найдено товарных строк.")
    return pd.DataFrame(rows)


def _significant_tokens(text: str) -> Set[str]:
    n = norm_txt(text)
    toks = set(re.findall(r"[a-zа-я0-9]+", n))
    return {t for t in toks if t not in _STOP_TOKENS and (len(t) > 2 or t.isdigit())}


def map_warehouse_to_store(warehouse: str, inventory_stores: Iterable[str]) -> Tuple[Optional[str], str]:
    """Map capitalization warehouse to a single inventory store.

    Returns (store_name_or_None, mapping_status).
    mapping_status: mapped | unmatched_store | uncertain_store
    """
    stores = [str(s).strip() for s in inventory_stores if str(s).strip()]
    if not warehouse or not stores:
        return None, STATUS_UNMATCHED_STORE

    wh = warehouse.strip()
    wh_low = wh.lower()

    # 1) exact (case-insensitive)
    exact = [s for s in stores if s.lower() == wh_low]
    if len(exact) == 1:
        return exact[0], "mapped"
    if len(exact) > 1:
        return None, STATUS_UNCERTAIN_STORE

    # 2) containment either way
    contained = [
        s for s in stores
        if s.lower() in wh_low or wh_low in s.lower()
    ]
    if len(contained) == 1:
        return contained[0], "mapped"
    if len(contained) > 1:
        return None, STATUS_UNCERTAIN_STORE

    # 3) token overlap (numbers / distinctive words)
    wh_toks = _significant_tokens(wh)
    scored: List[Tuple[int, str]] = []
    for s in stores:
        st = _significant_tokens(s)
        inter = wh_toks & st
        if not inter:
            continue
        # Prefer digit matches (store numbers)
        score = sum(10 if t.isdigit() else 3 for t in inter)
        scored.append((score, s))
    scored.sort(key=lambda x: (-x[0], x[1]))
    if not scored:
        return None, STATUS_UNMATCHED_STORE
    best_score = scored[0][0]
    top = [s for sc, s in scored if sc == best_score]
    if len(top) == 1 and best_score >= 10:
        return top[0], "mapped"
    if len(top) == 1 and best_score >= 3:
        # weak textual match — still unique
        return top[0], "mapped"
    return None, STATUS_UNCERTAIN_STORE


def _build_inventory_indexes(inv_df: pd.DataFrame) -> Dict[str, Dict[str, pd.DataFrame]]:
    """store -> { 'by_exact': dict name->df, 'by_key': dict sku_key->df } aggregated views."""
    indexes: Dict[str, Dict[str, pd.DataFrame]] = {}
    if inv_df.empty:
        return indexes
    work = inv_df.copy()
    if "sku_key" not in work.columns:
        work["sku_key"] = work["наименование"].map(sku_key)
    work["name_exact"] = work["наименование"].astype(str).str.strip()
    for store, grp in work.groupby("магазин"):
        indexes[str(store)] = {
            "by_exact": {n: g for n, g in grp.groupby("name_exact", dropna=False)},
            "by_key": {k: g for k, g in grp.groupby("sku_key", dropna=False) if k},
        }
    return indexes


def match_capitalization_to_inventory(
    cap_df: pd.DataFrame,
    inventory_df: pd.DataFrame,
) -> pd.DataFrame:
    """Exact-name match within mapped store only. Never cross stores."""
    stores = sorted(inventory_df["магазин"].dropna().astype(str).unique()) if not inventory_df.empty else []
    indexes = _build_inventory_indexes(inventory_df)

    # cache warehouse mapping
    wh_map: Dict[str, Tuple[Optional[str], str]] = {}
    out_rows = []

    for _, row in cap_df.iterrows():
        wh = str(row.get("склад", "") or "").strip()
        name = str(row.get("наименование", "") or "").strip()
        qty = float(row.get("оп_количество", 0) or 0)
        sm = float(row.get("оп_сумма", 0) or 0)
        key = str(row.get("sku_key") or sku_key(name))

        if wh not in wh_map:
            wh_map[wh] = map_warehouse_to_store(wh, stores)
        store, map_status = wh_map[wh]

        comment = ""
        linked_shortage = False
        inv_surplus = 0.0
        inv_shortage = 0.0
        match_docs = ""

        if map_status == STATUS_UNMATCHED_STORE:
            status = STATUS_UNMATCHED_STORE
            comment = "Склад оприходования не сопоставлен ни с одним магазином инвентаризации"
        elif map_status == STATUS_UNCERTAIN_STORE:
            status = STATUS_UNCERTAIN_STORE
            comment = "Неоднозначное сопоставление склада — данные не смешивались"
        else:
            assert store is not None
            idx = indexes.get(store, {})
            hit = idx.get("by_exact", {}).get(name)
            if hit is None or hit.empty:
                hit = idx.get("by_key", {}).get(key)
            if hit is None or hit.empty:
                status = STATUS_UNMATCHED_NAME
                comment = f"Наименование не найдено в инвентаризации склада «{store}» (точное совпадение)"
            else:
                inv_surplus = float(hit["излишек_сумма"].sum()) if "излишек_сумма" in hit.columns else 0.0
                inv_shortage = float(hit["недостача_сумма"].sum()) if "недостача_сумма" in hit.columns else 0.0
                shortage_qty = float(hit["недостача_кол"].sum()) if "недостача_кол" in hit.columns else 0.0
                linked_shortage = inv_shortage > 0 or shortage_qty > 0
                if "документ" in hit.columns:
                    match_docs = "; ".join(sorted(set(hit["документ"].astype(str).head(5))))
                if linked_shortage:
                    status = STATUS_MATCHED_SHORTAGE
                    comment = "Сопоставлено; по этой позиции в инвентаризации есть недостача"
                else:
                    status = STATUS_MATCHED
                    comment = "Сопоставлено по точному наименованию внутри склада"

        out_rows.append({
            "склад": wh,
            "магазин": store or "",
            "store_map_status": map_status if map_status != "mapped" else "mapped",
            "наименование": name,
            "оп_количество": qty,
            "оп_сумма": sm,
            "статус": status,
            "связь_с_недостачей": "ДА" if linked_shortage else "НЕТ",
            "инв_излишек_сумма": inv_surplus,
            "инв_недостача_сумма": inv_shortage,
            "документы": match_docs,
            "комментарий": comment,
            "sku_key": key,
        })

    result = pd.DataFrame(out_rows)
    if result.empty:
        return result
    return result.sort_values(["склад", "статус", "оп_сумма"], ascending=[True, True, False])


def capitalization_summary(matched_df: pd.DataFrame) -> Dict[str, float]:
    if matched_df is None or matched_df.empty:
        return {
            "cap_rows": 0,
            "cap_qty": 0.0,
            "cap_sum": 0.0,
            "matched": 0,
            "matched_shortage": 0,
            "unmatched_name": 0,
            "unmatched_store": 0,
            "uncertain_store": 0,
            "warehouses": 0,
            "matched_sum": 0.0,
            "shortage_linked_sum": 0.0,
        }
    st = matched_df["статус"]
    return {
        "cap_rows": int(len(matched_df)),
        "cap_qty": float(matched_df["оп_количество"].sum()),
        "cap_sum": float(matched_df["оп_сумма"].sum()),
        "matched": int((st == STATUS_MATCHED).sum()),
        "matched_shortage": int((st == STATUS_MATCHED_SHORTAGE).sum()),
        "unmatched_name": int((st == STATUS_UNMATCHED_NAME).sum()),
        "unmatched_store": int((st == STATUS_UNMATCHED_STORE).sum()),
        "uncertain_store": int((st == STATUS_UNCERTAIN_STORE).sum()),
        "warehouses": int(matched_df["склад"].nunique()),
        "matched_sum": float(
            matched_df.loc[st.isin([STATUS_MATCHED, STATUS_MATCHED_SHORTAGE]), "оп_сумма"].sum()
        ),
        "shortage_linked_sum": float(
            matched_df.loc[st == STATUS_MATCHED_SHORTAGE, "оп_сумма"].sum()
        ),
    }


def enrich_store_metrics_with_capitalization(
    store_metrics: pd.DataFrame,
    matched_df: pd.DataFrame,
) -> pd.DataFrame:
    """Attach capitalization totals to store ranking (by mapped inventory store)."""
    sm = store_metrics.copy()
    if sm.empty:
        return sm
    sm["оприходование_сумма"] = 0.0
    sm["оприходование_кол"] = 0.0
    sm["оприходование_строк"] = 0
    sm["оп_связь_с_недостачей"] = 0
    if matched_df is None or matched_df.empty:
        return sm

    usable = matched_df[matched_df["магазин"].astype(str).str.len() > 0]
    if usable.empty:
        return sm
    agg = usable.groupby("магазин").agg(
        оприходование_сумма=("оп_сумма", "sum"),
        оприходование_кол=("оп_количество", "sum"),
        оприходование_строк=("наименование", "count"),
        оп_связь_с_недостачей=("связь_с_недостачей", lambda s: int((s == "ДА").sum())),
    )
    sm = sm.set_index("магазин")
    for col in agg.columns:
        sm[col] = agg[col]
    sm[list(agg.columns)] = sm[list(agg.columns)].fillna(0)
    return sm.reset_index()
