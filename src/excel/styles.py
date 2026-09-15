# -*- coding: utf-8 -*-
"""Excel visual styles and low-level cell helpers (preserved from v3.0)."""
from __future__ import annotations

import re
from typing import Dict, Optional

from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ACCENT = "1F4E79"
ACCENT2 = "2E75B6"
GD = "375623"
RD = "C00000"
ORG = "BF8F00"
LB = "DEEAF1"
GF = "E2EFDA"
RF = "FFE0E0"
YF = "FFF2CC"
WH = "FFFFFF"
HDR_F = Font(name="Calibri", bold=True, color=WH, size=10)
SUBT_F = Font(name="Calibri", bold=True, size=10, color="404040")
TITL_F = Font(name="Calibri", bold=True, size=14, color=ACCENT)
HD = PatternFill("solid", fgColor=ACCENT)
HM = PatternFill("solid", fgColor=ACCENT2)
HG = PatternFill("solid", fgColor=GD)
HR = PatternFill("solid", fgColor=RD)
SF = PatternFill("solid", fgColor=GF)
DF = PatternFill("solid", fgColor=RF)
PF = PatternFill("solid", fgColor=YF)
TF = PatternFill("solid", fgColor=LB)
SH = PatternFill("solid", fgColor="D9E2F3")
_th = Side(style="thin", color="BFBFBF")
BD = Border(left=_th, right=_th, top=_th, bottom=_th)
NUM = "#,##0.000"
SUM = "#,##0.00"
INT = "#,##0"
PCT = "0.0%"

SCENARIO_COLORS = {
    "СЦ-1": "FFF2CC", "СЦ-2А": "E2EFDA", "СЦ-2Б": "E2EFDA", "СЦ-2В": "E2EFDA",
    "СЦ-3": "FFCCCC", "СЦ-4": "FFE0E0", "СЦ-5": "FDE9D9", "СЦ-6": "F2F2F2", "СЦ-7": "DEEAF1",
}
SCENARIO_PRIORITY_COLOR = {1: "C00000", 2: "FF0000", 3: "BF8F00", 4: "2E75B6", 5: "7F7F7F"}

LEVEL_RU = {
    "intra_doc": "Внутри документа",
    "cross_doc": "Между документами магазина",
    "cross_store_analytical": "Между магазинами (аналитика)",
}

HOMOGENEITY_RU = {
    "same_category": "Одна категория",
    "same_article": "Один артикул",
    "product_key": "Схожий товар",
    "fallback_product_key": "Резервная группа (проверить)",
    "rejected_category": "Разные категории",
}

def GL(c: int) -> str:
    return get_column_letter(c)


def mtitle(ws, row, c1, c2, text, font=None, fill=None, align="center"):
    ws.merge_cells(f"{GL(c1)}{row}:{GL(c2)}{row}")
    c = ws.cell(row=row, column=c1, value=text)
    c.font = font or TITL_F
    c.fill = fill or PatternFill()
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True,
                            indent=1 if align == "left" else 0)
    c.border = BD
    return c


def hrow(ws, row, headers, fill=None, ht=26, col_start=1):
    fills = fill if isinstance(fill, list) else [fill or HD] * len(headers)
    for i, h in enumerate(headers):
        col = col_start + i
        c = ws.cell(row=row, column=col, value=h)
        c.font = HDR_F
        c.fill = fills[i]
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BD
    ws.row_dimensions[row].height = ht


def zebra_fill(row_idx: int) -> Optional[PatternFill]:
    return PatternFill("solid", fgColor="FAFAFA") if row_idx % 2 else None


LEVEL_RU = {
    "intra_doc": "Внутри документа",
    "cross_doc": "Между документами магазина",
    "cross_store_analytical": "Между магазинами (аналитика)",
}

HOMOGENEITY_RU = {
    "same_category": "Одна категория",
    "same_article": "Один артикул",
    "product_key": "Схожий товар",
    "fallback_product_key": "Резервная группа (проверить)",
    "rejected_category": "Разные категории",
}


def tc(ws, row, col, val, fill=None, bold=False, wrap=False, fmt=None):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(name="Calibri", size=10, bold=bold)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=wrap, indent=1)
    c.border = BD
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def nc(ws, row, col, val, fmt=SUM, fill=None, bold=False):
    c = ws.cell(row=row, column=col, value=val if val is not None else 0)
    c.number_format = fmt
    c.font = Font(name="Calibri", size=10, bold=bold)
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.border = BD
    if fill:
        c.fill = fill
    return c


def trow(ws, row, nc_, texts, nums, ht=18):
    for c in range(1, nc_ + 1):
        ws.cell(row=row, column=c).fill = TF
        ws.cell(row=row, column=c).border = BD
    for col, val in texts.items():
        tc(ws, row, col, val, TF, bold=True)
    for col, (val, fmt) in nums.items():
        nc(ws, row, col, val, fmt, TF, bold=True)
    ws.row_dimensions[row].height = ht


def dbars(ws, rng, color="4472C4"):
    if not rng or ":" not in str(rng):
        return
    try:
        l_row = int(re.findall(r"\d+", str(rng).split(":")[0])[0])
        r_row = int(re.findall(r"\d+", str(rng).split(":")[1])[0])
        if r_row < l_row:
            return
    except Exception:
        return
    ws.conditional_formatting.add(str(rng), DataBarRule(start_type="min", end_type="max", color=color))


def cs3(ws, rng):
    if not rng or ":" not in str(rng):
        return
    try:
        l_row = int(re.findall(r"\d+", str(rng).split(":")[0])[0])
        r_row = int(re.findall(r"\d+", str(rng).split(":")[1])[0])
        if r_row < l_row:
            return
    except Exception:
        return
    ws.conditional_formatting.add(
        str(rng),
        ColorScaleRule(
            start_type="min", start_color="63BE7B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="F8696B",
        ),
    )


def safe_sheet_name(name: str, used: set) -> str:
    s = re.sub(r'[\\/*?:\[\]]', "_", str(name).strip())[:28]
    if not s:
        s = "Лист"
    base = s
    n = 1
    while s in used:
        suffix = f"_{n}"
        s = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(s)
    return s


def set_col_widths(ws, widths: Dict[int, float]):
    for col, w in widths.items():
        ws.column_dimensions[GL(col)].width = w
